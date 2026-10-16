"""Builds report.docx from outputs/ - run main.py (w/ --compare) and score.py first."""
import json
from datetime import datetime

import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt

from src import config as C, plots
from src.data import clean, load
from src.trend import LevelTrend

FIG = C.OUT / "figures"


def save(fig, name):
    FIG.mkdir(parents=True, exist_ok=True)
    path = FIG / f"{name}.png"
    fig.savefig(path)
    plt.close(fig)
    return path


def table(doc, df, fmt=None):
    fmt = fmt or {}
    t = doc.add_table(rows=1, cols=len(df.columns))
    t.style = "Light Grid Accent 1"
    for i, c in enumerate(df.columns):
        t.rows[0].cells[i].text = str(c)
    for _, r in df.iterrows():
        cells = t.add_row().cells
        for i, c in enumerate(df.columns):
            v = r[c]
            cells[i].text = fmt[c].format(v) if c in fmt and pd.notna(v) else str(v)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
    return t


def main():
    doc = Document()
    doc.styles["Normal"].font.size = Pt(10.5)
    props = doc.core_properties
    props.author = props.last_modified_by = "Almila Duru Kavak"
    props.title = "Freight rate prediction - Spotter take-home"
    props.comments = ""
    props.created = props.modified = datetime.now()
    tr = clean(load(C.TRAIN_CSV))
    dq = pd.read_csv(C.OUT / "data_quality.csv")
    comp = pd.read_csv(C.OUT / "model_comparison.csv")
    ch = pd.read_csv(C.OUT / "city_holdout.csv")
    fi = pd.read_csv(C.OUT / "feature_importance.csv")
    m = json.loads((C.OUT / "metrics.json").read_text())
    dec = pd.read_csv(C.DEC_PRED_CSV)

    doc.add_heading("Freight rate prediction - Spotter take-home", 0)
    doc.add_paragraph(
        "Goal: predict posted_rate for 12,000 validation loads (Nov-Dec 2025) from Jan-Oct 2025 history, "
        "plus a daily December 2025 rate for Lexington -> Fort Wayne (360 mi, Dry Van, 32,000 lb). "
        "No metric was given, so I optimise MAE/MAPE and report RMSE too.")

    doc.add_heading("1. Data quality", 1)
    doc.add_paragraph(
        "Every check below is computed by src/data.py:audit and written to outputs/data_quality.csv on each run.")
    table(doc, dq[dq["issue"].ne("rows") & dq["issue"].ne("cities")][["file", "issue", "count", "action"]])
    doc.add_paragraph(
        "Negative weights: the absolute values have the same distribution (5k-47.5k) and the same rate per mile "
        "as positive rows, so it's a sign glitch and abs() is safe. Missing weight is imputed with the train "
        "median per equipment. Rate outliers form two tight clusters (~0.3x and ~3.5x the lane norm) - label noise, "
        "not market moves. They are kept in training (robust loss handles them, see table) and never removed from "
        "any evaluation set.")
    doc.add_picture(str(save(plots.outliers(tr), "outliers")), width=Inches(6))

    doc.add_heading("2. Key EDA findings", 1)
    for s in [
        "Rate is mostly distance x rate-per-mile; rpm falls w/ distance and Reefer > Flatbed > Dry Van. "
        "Target = log(rate / distance), predictions are multiplied back by distance.",
        "Market level over time is a hump (-4% in Jan, +4% in Jun, -1% in Aug, back to ~+0.7% by Oct), not a steady drift.",
        "market_index has the same rough shape but a much larger swing and doesn't line up out of time "
        "(Sep has the lowest index but a normal rate level). Adding it made 3 of 4 validation windows worse (holdout MAE 127 vs 85).",
        "8 cities (12% of validation rows) never appear in train, so the model can't lean on city IDs alone; "
        "lat/lon + distance carry new lanes.",
    ]:
        doc.add_paragraph(s, style="List Bullet")
    doc.add_picture(str(save(plots.rpm_vs_distance(tr), "rpm_distance")), width=Inches(6))
    doc.add_picture(str(save(plots.monthly_trend(tr), "monthly_trend")), width=Inches(6.5))

    doc.add_heading("3. Validation approach", 1)
    doc.add_paragraph(
        "Validation is Nov-Dec, so everything is evaluated forward in time - no random splits. "
        "Main holdout: train on Jan-Aug, score Sep-Oct (same 1-2 month horizon as the real task). "
        "Plus 3 expanding-window folds (train up to Jul/Aug/Sep, score the next month). "
        "All encoders, imputation medians, lane medians and the trend are fit inside each training window only. "
        "Headline metrics are on ALL rows; 'clean' (excluding the label-noise outliers) is secondary. "
        "Seen vs unseen lanes are reported separately, and a city-holdout sim hides 8 random cities from training "
        "(3 repeats) to mimic the cities that only appear in validation. After choosing the config, the model is "
        "refit on all Jan-Oct data for the final predictions.")

    doc.add_heading("4. Model comparison", 1)
    allr = comp[comp.subset == "all"].copy()
    clean_ = comp[comp.subset == "clean"].set_index("model")["ho_MAPE"]
    allr["ho_MAPE_clean"] = allr["model"].map(clean_)
    allr = allr.sort_values("ho_MAE")
    cols = ["model", "cv_MAE", "cv_MAPE", "ho_MAE", "ho_RMSE", "ho_MAPE", "ho_MAPE_clean"]
    table(doc, allr[cols], {c: "{:.1f}" for c in ["cv_MAE", "ho_MAE", "ho_RMSE"]} | {c: "{:.2f}%" for c in ["cv_MAPE", "ho_MAPE", "ho_MAPE_clean"]})
    doc.add_paragraph("cv_* = mean over the 3 monthly folds, ho_* = Sep-Oct holdout. RMSE is dominated by the "
                      "~1.4% outlier labels, which no model can predict.")
    sl = comp[(comp.model == "HGB final") & comp.subset.isin(["all", "seen lane", "unseen lane"])]
    doc.add_paragraph("Final model by lane type (holdout):")
    table(doc, sl[["subset", "ho_n", "ho_MAE", "ho_MAPE"]], {"ho_n": "{:.0f}", "ho_MAE": "{:.1f}", "ho_MAPE": "{:.2f}%"})
    doc.add_paragraph("Unseen-city sim (8 cities hidden from train, rows touching them, mean of 3 repeats; the lane baseline has no "
                      "history for those lanes so it falls back to global rpm x distance):")
    chs = ch.groupby("model")[["n", "MAE", "MAPE"]].mean().reset_index()
    table(doc, chs, {"n": "{:.0f}", "MAE": "{:.1f}", "MAPE": "{:.2f}%"})

    doc.add_heading("5. Final model", 1)
    for s in FINAL_NOTES:
        doc.add_paragraph(s, style="List Bullet")
    doc.add_picture(str(save(plots.trend_fit(LevelTrend("level").fit(tr), "Market level - tested trend component (not in final model)"), "trend")), width=Inches(6))
    doc.add_paragraph("Feature importance - permutation on the Sep-Oct holdout (model trained Jan-Aug):")
    doc.add_picture(str(save(plots.feature_importance(fi), "importance")), width=Inches(5.5))

    doc.add_heading("6. December chart (score.py)", 1)
    doc.add_paragraph(
        f"Lexington -> Fort Wayne, 360 mi, Dry Van, 32,000 lb. Predictions range "
        f"${dec.predicted_rate.min():.0f}-${dec.predicted_rate.max():.0f} "
        f"(${dec.predicted_rate.min() / 360:.2f}-${dec.predicted_rate.max() / 360:.2f}/mi). "
        "City coords come from the train lookup. quote_signal isn't in the chart inputs, so each December date is "
        "predicted w/ every real quote_signal value from validation.csv loads on that date, then averaged - i.e. the "
        "expected rate for a load that day. Plugging in the daily median instead came out ~5% low: on this lane the "
        "model's quote_signal curve has its dip right at the median. Day-to-day movement comes from day of week + quote_signal; there's no time trend, "
        "so there's no extrapolated slope or made-up seasonality. For reference, the 21 Dry Van loads on this lane in train "
        "averaged $823 (median $808, range $758-$934, Jan-Oct).")
    doc.add_picture(str(C.ROOT / "scorer_results" / "candidate_december.png"), width=Inches(6.5))

    doc.add_heading("7. Caveats", 1)
    for s in CAVEATS:
        doc.add_paragraph(s, style="List Bullet")
    doc.save(C.ROOT / "report.docx")
    print("wrote report.docx")


FINAL_NOTES = [
    "sklearn HistGradientBoostingRegressor (installs cleanly on Python 3.14/macOS w/ pip, handles NaN natively; "
    "LightGBM/XGBoost need libomp on macOS so I skipped them).",
    "Target: log(rate per mile); prediction = exp(trees) x distance. "
    "The rpm target removes most of the distance scale, the log makes errors relative.",
    "Absolute-error loss: robust to the ~1.4% noisy labels. With it, dropping or down-weighting outliers "
    "made no measurable difference, so nothing is dropped. Squared loss was clearly worse.",
    "Features: lat/lon, distance, equipment, weight, quote_signal, day of week. No city IDs - coords-only ties "
    "on CV, wins the holdout (85.3 vs 86.8 MAE) and treats new cities exactly like known ones. No month/day-of-year (Nov/Dec never seen in train).",
    "market_index left out: hurt every window except Aug (holdout MAE 127 vs 85). Its relation to rates isn't stable over time.",
    "No time trend in the final model. Tested: time index as a tree feature (flat extrapolation, clearly worse), "
    "linear/damped slopes on the detrended target (early scratch run, overshot - the level is a hump, not a drift) "
    "and a market-level "
    "component (smoothed daily rate vs lane norm, held flat at the last-60-day level). The level component only won "
    "the Aug fold, right after the Jun peak, and lost Sep, Oct and the Sep-Oct holdout, so it's off. The model "
    "therefore predicts at the Jan-Oct average level, which is close to the Aug-Oct level (~0-1% apart).",
    "Fallback chain: model -> lane median rpm x distance -> global rpm x distance, w/ asserts for no NaN / "
    "non-positive values. The model produced every final prediction (fallback never triggered).",
]
CAVEATS = [
    "Weak spot: in the Aug fold (train through Jul, right after the Jun peak) the model was worse than the lane-median "
    "baseline (MAE 174 vs 159) - it lags when the market level drops fast. The level-trend variant handled that fold "
    "best (155) but lost everywhere else. If Nov-Dec has a sharp level move, expect a similar lag.",
    "No December history in train, so holiday seasonality can't be learned; the December line is the average level + weekday pattern.",
    "Only 21 holdout rows are on lanes unseen in Jan-Aug, so the seen/unseen lane split is thin; the city-holdout sim is the real test there.",
    "No metric was specified; MAE/MAPE are optimised. If Spotter scores on RMSE, the outlier labels will dominate.",
    "quote_signal is used as given in validation.csv; if it's not known at quote time in production that's leakage "
    "risk worth checking. For the December chart I averaged over the real values seen each day in validation.csv.",
    "Coordinates look synthetic/clipped (Laredo lat 25.5, Boston/Providence lon -69.5); fine here since every "
    "city has one consistent coord pair, but worth knowing.",
]

if __name__ == "__main__":
    main()
